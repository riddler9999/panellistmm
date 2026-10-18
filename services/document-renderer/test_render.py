import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image
from render import render


def contract(kind, **kwargs):
    return {"contract": {"document_type": kind, "contract_version": "2.0", "status": "ready", "title": "ဝန်ထမ်း စီမံခန့်ခွဲမှု", "review_required": True, "assumptions": [], **kwargs}}


ORG = contract("org_chart", nodes=[
    {"id": "ceo", "label": "အုပ်ချုပ်မှု ဒါရိုက်တာ", "person": "ဦးအောင်", "parent_id": None},
    {"id": "hr", "label": "လူ့စွမ်းအားအရင်းအမြစ်", "parent_id": "ceo"},
    {"id": "ops", "label": "=HYPERLINK(\"https://example.com\")", "parent_id": "ceo"},
])
SOP = contract("sop", steps=[
    {"step_no": 1, "step_type": "task", "role": "ဝန်ထမ်း", "action": "ခွင့်တင်ပါ", "next_steps": [2]},
    {"step_no": 2, "step_type": "decision", "role": "မန်နေဂျာ", "action": "ခွင့်ပြုမလား", "next_steps": [3, "END"], "branches": [{"condition": "ခွင့်ပြု", "next_step": 3}, {"condition": "ခွင့်မပြု", "next_step": "END"}]},
    {"step_no": 3, "step_type": "task", "role": "HR", "action": "မှတ်တမ်းတင်ပါ", "next_steps": ["END"]},
])


class RendererTests(unittest.TestCase):
    def test_both_outputs_and_literal_formulas(self):
        with tempfile.TemporaryDirectory() as directory:
            for payload in (ORG, SOP):
                xlsx, png = Path(directory) / "result.xlsx", Path(directory) / "result.png"
                render(payload, xlsx, png)
                with Image.open(png) as image:
                    self.assertEqual(image.format, "PNG")
                    image.verify()
                workbook = load_workbook(xlsx)
                self.assertIn("Diagram preview", workbook.sheetnames)
                self.assertTrue(all(c.data_type != "f" for s in workbook for row in s for c in row))

    def test_bad_graphs_fail_without_outputs(self):
        invalid = [contract("org_chart", nodes=[{"id": "x", "label": "x", "parent_id": "missing"}]), contract("sop", steps=[{"step_no": 1, "action": "Loop", "next_steps": [1]}])]
        with tempfile.TemporaryDirectory() as directory:
            for payload in invalid:
                with self.assertRaises(ValueError):
                    render(payload, Path(directory) / "result.xlsx", Path(directory) / "result.png")
            self.assertEqual(list(Path(directory).iterdir()), [])


if __name__ == "__main__":
    unittest.main()
