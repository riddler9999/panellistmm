"""Deterministic XLSX/PNG rendering of reviewed semantic document contracts."""
import json
import math
import os
import re
import sys
import unicodedata
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as SheetImage
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageDraw, ImageFont, features


def text(value):
    return str(value or "")[:2000]


def graph(contract):
    kind = contract.get("document_type")
    if kind == "org_chart":
        nodes = contract.get("nodes", [])
        if not 1 <= len(nodes) <= 100:
            raise ValueError("ORG_NODE_LIMIT")
        by = {n["id"]: n for n in nodes}
        if len(by) != len(nodes):
            raise ValueError("DUPLICATE_ID")
        roots = [n for n in nodes if n.get("parent_id") is None]
        if len(roots) != 1:
            raise ValueError("SINGLE_ROOT_REQUIRED")
        children = {key: [] for key in by}
        for n in nodes:
            parent = n.get("parent_id")
            if parent is not None:
                if parent not in by:
                    raise ValueError("MISSING_PARENT")
                children[parent].append(n["id"])
        positions, seen, leaf = {}, set(), [0]

        def walk(key, depth):
            if key in seen:
                raise ValueError("ORG_CYCLE")
            seen.add(key)
            for child in children[key]:
                walk(child, depth + 1)
            x = sum(positions[k][0] for k in children[key]) / len(children[key]) if children[key] else leaf[0]
            if not children[key]:
                leaf[0] += 1
            positions[key] = (x, depth)

        walk(roots[0]["id"], 0)
        if len(seen) != len(nodes):
            raise ValueError("DISCONNECTED_OR_CYCLIC_ORG")
        labels = {n["id"]: "\n".join(text(n.get(k)) for k in ("label", "person", "department") if n.get(k)) for n in nodes}
        edges = [(n["parent_id"], n["id"], "") for n in nodes if n.get("parent_id") is not None]
        headers = ["ID", "Reports to", "Role", "Person", "Department", "Type"]
        rows = [[n.get(k, "") for k in ("id", "parent_id", "label", "person", "department", "type")] for n in nodes]
    elif kind == "sop":
        nodes = contract.get("steps", [])
        if not 1 <= len(nodes) <= 100:
            raise ValueError("SOP_STEP_LIMIT")
        by = {n["step_no"]: n for n in nodes}
        if len(by) != len(nodes):
            raise ValueError("DUPLICATE_STEP")
        nodes = sorted(nodes, key=lambda n: n["step_no"])
        positions = {n["step_no"]: (0, i) for i, n in enumerate(nodes)}
        positions["END"] = (0, len(nodes))
        labels = {n["step_no"]: f'{n["step_no"]}. {text(n.get("action"))}\n{text(n.get("role"))}' for n in nodes}
        labels["END"] = "END"
        edges = []
        for n in nodes:
            targets = n.get("next_steps", [])
            if not targets or any(t != "END" and t not in by for t in targets):
                raise ValueError("INVALID_SOP_TARGET")
            branches = n.get("branches", [])
            if n.get("step_type") == "decision":
                if len(branches) < 2 or {b.get("next_step", b.get("target")) for b in branches} != set(targets):
                    raise ValueError("INVALID_BRANCHES")
                edges.extend((n["step_no"], b.get("next_step", b.get("target")), text(b.get("condition"))) for b in branches)
            else:
                edges.extend((n["step_no"], t, "") for t in targets)
        reachable = {nodes[0]["step_no"]}
        can_end = {"END"}
        for _ in range(len(nodes) + 1):
            reachable.update(b for a, b, _ in edges if a in reachable)
            can_end.update(a for a, b, _ in edges if b in can_end)
        if any(k not in reachable or k not in can_end for k in by):
            raise ValueError("UNREACHABLE_OR_NONTERMINATING_SOP")
        headers = ["Step", "Type", "Action", "Role", "Department", "Decision", "Next", "Branches"]
        rows = [[n.get(k, "") for k in ("step_no", "step_type", "action", "role", "department", "decision")] + [json.dumps(n.get("next_steps", []), ensure_ascii=False), json.dumps(n.get("branches", []), ensure_ascii=False)] for n in nodes]
    else:
        raise ValueError("UNSUPPORTED_DOCUMENT_TYPE")
    if any(not label.strip() for label in labels.values()):
        raise ValueError("EMPTY_LABEL")
    return kind, positions, labels, edges, headers, rows


def render(payload, xlsx_path, png_path):
    c = payload["contract"]
    if c.get("status") != "ready" or c.get("contract_version") != "2.0":
        raise ValueError("READY_V2_CONTRACT_REQUIRED")
    kind, positions, labels, edges, headers, rows = graph(c)
    if not features.check("raqm"):
        raise RuntimeError("MYANMAR_SHAPING_UNAVAILABLE")
    font_path = os.environ.get("PANELLIST_FONT", "/usr/share/fonts/truetype/noto/NotoSansMyanmar-Regular.ttf")
    font = ImageFont.truetype(font_path, 22)
    small = ImageFont.truetype(font_path, 18)
    latin_path = os.environ.get("PANELLIST_LATIN_FONT", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    latin = ImageFont.truetype(latin_path, 22)
    latin_small = ImageFont.truetype(latin_path, 18)
    scratch = ImageDraw.Draw(Image.new("RGB", (1, 1)))

    def runs(label, compact=False):
        for run in re.findall(r"[\u1000-\u109f\uaa60-\uaa7f\ua9e0-\ua9ff]+|[^\u1000-\u109f\uaa60-\uaa7f\ua9e0-\ua9ff]+", label):
            myanmar = bool(re.match(r"[\u1000-\u109f\uaa60-\uaa7f\ua9e0-\ua9ff]", run))
            yield run, (small if compact else font) if myanmar else (latin_small if compact else latin)

    def measure(label):
        return sum(scratch.textlength(run, font=face) for run, face in runs(label))

    def draw_text(at, label, fill, compact=False):
        x, y = at
        for run, face in runs(label, compact):
            draw.text((x, y + 30), run, fill=fill, font=face, anchor="ls")
            x += draw.textlength(run, font=face)

    def wrap(label, width):
        lines = []
        for paragraph in label.splitlines():
            current = ""
            clusters = []
            for char in paragraph:
                if clusters and (unicodedata.category(char).startswith("M") or clusters[-1].endswith("\u1039")):
                    clusters[-1] += char
                else:
                    clusters.append(char)
            for char in clusters:
                if current and measure(current + char) > width:
                    lines.append(current)
                    current = ""
                current += char
            lines.append(current)
        return lines

    lines = {key: wrap(value, 330) for key, value in labels.items()}
    box_w, box_h = 380, max(150, max(len(v) for v in lines.values()) * 42 + 30)
    gap_x, gap_y = 80, 100
    width = math.ceil((max(x for x, y in positions.values()) + 1) * (box_w + gap_x) + (420 if kind == "sop" else 60))
    height = math.ceil((max(y for x, y in positions.values()) + 1) * (box_h + gap_y) + 150)
    if width * height > 32_000_000 or max(width, height) > 30000:
        raise ValueError("DIAGRAM_SIZE_LIMIT")
    image = Image.new("RGB", (width, height), "#f8fafc")
    draw = ImageDraw.Draw(image)
    draw_text((30, 20), text(c.get("title"))[:120], "#172554")
    coords = {key: (round(30 + x * (box_w + gap_x)), round(110 + y * (box_h + gap_y))) for key, (x, y) in positions.items()}
    for index, (source, target, condition) in enumerate(edges):
        x1, y1 = coords[source]
        x2, y2 = coords[target]
        start, end = (x1 + box_w // 2, y1 + box_h), (x2 + box_w // 2, y2)
        if kind == "sop" and (condition or positions[target][1] != positions[source][1] + 1):
            lane = box_w + 100 + (index % 4) * 65
            start, end = (x1 + box_w, y1 + box_h // 2), (x2 + box_w, y2 + box_h // 2)
            draw.line([start, (lane, start[1]), (lane, end[1]), end], fill="#64748b", width=3)
            draw.polygon([end, (end[0] + 12, end[1] - 7), (end[0] + 12, end[1] + 7)], fill="#64748b")
            for line_index, line in enumerate(wrap(condition, 300)):
                draw_text((lane + 8, start[1] + line_index * 34), line, "#334155", compact=True)
        else:
            middle = (start[1] + end[1]) // 2
            draw.line([start, (start[0], middle), (end[0], middle), end], fill="#64748b", width=3)
            draw.polygon([end, (end[0] - 7, end[1] - 12), (end[0] + 7, end[1] - 12)], fill="#64748b")
    for key, (x, y) in coords.items():
        draw.rounded_rectangle((x, y, x + box_w, y + box_h), radius=16, fill="#ffffff", outline="#7c3aed", width=3)
        for index, line in enumerate(lines[key]):
            draw_text((x + 22, y + 12 + index * 42), line, "#172554")
    image.save(png_path)
    book = Workbook()
    sheet = book.active
    sheet.title = "Organization" if kind == "org_chart" else "SOP Steps"
    for row in [headers] + rows:
        sheet.append([text(value) for value in row])
    for row in sheet:
        for cell in row:
            cell.data_type = "s"  # User-supplied formula strings remain literal text.
            cell.font = Font(name="Noto Sans Myanmar", size=11, bold=cell.row == 1)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor="EDE9FE")
        sheet.row_dimensions[row[0].row].height = 70
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 32
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    meta = book.create_sheet("Read me")
    for row in [["Title", text(c.get("title"))], ["Review required", str(c.get("review_required", True))], ["Editing", "Edit the data sheet. PNG is a snapshot and does not update automatically."], ["Assumptions", "\n".join(map(text, c.get("assumptions", [])))]]:
        meta.append(row)
    for row in meta:
        for cell in row:
            cell.data_type = "s"
    meta.column_dimensions["B"].width = 100
    preview = book.create_sheet("Diagram preview")
    embedded = SheetImage(str(png_path))
    scale = min(1, 1600 / width)
    embedded.width, embedded.height = width * scale, height * scale
    preview.add_image(embedded, "A1")
    book.save(xlsx_path)


if __name__ == "__main__":
    source = Path(sys.argv[1])
    if source.stat().st_size > 500_000:
        raise ValueError("INPUT_SIZE_LIMIT")
    render(json.loads(source.read_text()), sys.argv[2], sys.argv[3])
